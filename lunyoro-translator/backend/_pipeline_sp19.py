"""
_pipeline_sp19.py
=================
Full pipeline for sentence pair 19:
  1. Extract + clean pairs from xlsx
  2. Merge into new_only_train/val (dedup against full train+val)
  3. Train MarianMT --new-only both directions  (5 epochs)
  4. Train NLLB     --new-only both directions  (3 epochs)
  5. Push all 4 models to HF Hub
"""
import re, shutil, subprocess, sys, unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

BASE         = Path(__file__).parent
RAW_DIR      = BASE / "data" / "raw"
CLEANED_DIR  = BASE / "data" / "cleaned"
TRAINING_DIR = BASE / "data" / "training"

INPUT_FILE = RAW_DIR / "sentence pair 19.xlsx"
CLEAN_FILE = CLEANED_DIR / "sentence_pair_19_clean.csv"
TRAIN_CSV  = TRAINING_DIR / "new_only_train.csv"
VAL_CSV    = TRAINING_DIR / "new_only_val.csv"

VENV_PY = BASE / "venv312" / "Scripts" / "python.exe"
PY = str(VENV_PY) if VENV_PY.exists() else sys.executable

# ─────────────────────────────────────────────────────────────────────────────
print("=" * 60)
print("  SP19 PIPELINE")
print(f"  Input: {INPUT_FILE.name}")
print("=" * 60)


def _norm(text):
    text = str(text).strip()
    text = re.sub(r"\s+", " ", text)
    return (unicodedata.normalize("NFC", text)
            .replace("\u2018", "'").replace("\u2019", "'")
            .replace("\u201c", '"').replace("\u201d", '"'))


# ── Step 1: Extract & clean ───────────────────────────────────────────────────
print(f"\n[1/5] Extracting and cleaning {INPUT_FILE.name} ...")

wb   = load_workbook(INPUT_FILE)
ws   = wb.active
rows = list(ws.iter_rows(values_only=True))
print(f"  Raw rows: {len(rows)}")

# Column layout (no header):
#   0=row_id  1=group_id  2=orig_en  3=orig_tense
#   4=orig_lun  5=target_tense  6=varied_en  7=varied_lun
orig   = pd.DataFrame([(r[2], r[4]) for r in rows if r[2] and r[4]], columns=["english", "lunyoro"])
varied = pd.DataFrame([(r[6], r[7]) for r in rows if r[6] and r[7]], columns=["english", "lunyoro"])
df = pd.concat([orig, varied], ignore_index=True)
print(f"  Combined pairs (original + varied): {len(df)}")

df["english"] = df["english"].apply(_norm)
df["lunyoro"] = df["lunyoro"].apply(_norm).str.lower()
df = df.dropna()
df = df[df["english"].str.strip().astype(bool) & df["lunyoro"].str.strip().astype(bool)]
df = df[df["english"].str.lower() != df["lunyoro"]]          # drop identity pairs
df = df[df["lunyoro"].str.split().str.len() >= 3]            # require ≥3 Lunyoro tokens
df = df.drop_duplicates(subset=["english", "lunyoro"])

CLEANED_DIR.mkdir(parents=True, exist_ok=True)
df.to_csv(CLEAN_FILE, index=False)
print(f"  Clean pairs: {len(df)}  →  {CLEAN_FILE.name}")


# ── Step 2: Merge into new_only_train / new_only_val ─────────────────────────
print(f"\n[2/5] Merging into training set (dedup against full train+val) ...")

# Dedup against the *full* corpus so we don't re-add anything already trained on
full_train = TRAINING_DIR / "train.csv"
full_val   = TRAINING_DIR / "val.csv"
existing   = set()
for csv_path in (full_train, full_val, TRAIN_CSV, VAL_CSV):
    if csv_path.exists():
        tmp = pd.read_csv(csv_path)
        existing |= set(zip(
            tmp["english"].str.lower().str.strip(),
            tmp["lunyoro"].str.lower().str.strip(),
        ))

new = df[~df.apply(
    lambda r: (r["english"].lower().strip(), r["lunyoro"].lower().strip()) in existing,
    axis=1,
)]
print(f"  Genuinely new pairs: {len(new)}")

if len(new) > 0:
    # Backup before modifying
    for p in (TRAIN_CSV, VAL_CSV):
        if p.exists():
            shutil.copy(p, str(p) + ".bak")

    split = int(len(new) * 0.9)
    new.iloc[:split][["english", "lunyoro"]].to_csv(TRAIN_CSV, mode="a", index=False, header=False)
    new.iloc[split:][["english", "lunyoro"]].to_csv(VAL_CSV,   mode="a", index=False, header=False)

    t2 = pd.read_csv(TRAIN_CSV)
    v2 = pd.read_csv(VAL_CSV)
    print(f"  new_only_train: {len(t2):,}  (+{len(new.iloc[:split])})")
    print(f"  new_only_val  : {len(v2):,}  (+{len(new.iloc[split:])})")

    # Timestamped incremental snapshot
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    new.to_csv(CLEANED_DIR / f"incremental_{ts}.csv", index=False)
    print(f"  Incremental snapshot: incremental_{ts}.csv")
else:
    print("  Already merged — nothing to do.")


# ── Step 3: MarianMT training ─────────────────────────────────────────────────
print(f"\n[3/5] Training MarianMT --new-only both directions (5 epochs) ...")
subprocess.run(
    [PY, "train_marian.py", "--new-only", "--direction", "both",
     "--epochs", "5", "--batch-size", "32", "--lr", "3e-5"],
    check=True,
)


# ── Step 4: NLLB training ─────────────────────────────────────────────────────
print(f"\n[4/5] Training NLLB --new-only both directions (3 epochs) ...")
subprocess.run(
    [PY, "train_nllb.py", "--new-only", "--direction", "both",
     "--epochs", "3", "--batch-size", "16", "--lr", "8e-6", "--fp16"],
    check=True,
)


# ── Step 5: Push to HF Hub ────────────────────────────────────────────────────
print(f"\n[5/5] Pushing all 4 models to HF Hub ...")
for model in ["en2lun", "lun2en", "nllb_en2lun", "nllb_lun2en"]:
    subprocess.run([PY, "push_models.py", "--model", model], check=True)

print("\n" + "=" * 60)
print("  SP19 PIPELINE COMPLETE")
print("=" * 60)
