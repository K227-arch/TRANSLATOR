"""
diagnose_sp19.py — show exactly which pairs were dropped from sentence pair 19.xlsx and why.
"""
import re, unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE = Path(__file__).parent

def _norm(text):
    text = str(text).strip()
    text = re.sub(r"\s+", " ", text)
    return (unicodedata.normalize("NFC", text)
            .replace("\u2018", "'").replace("\u2019", "'")
            .replace("\u201c", '"').replace("\u201d", '"'))

f = BASE / "data" / "raw" / "sentence pair 19.xlsx"
wb = load_workbook(f)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))

# Build full raw pair list
all_pairs = []
for r in rows:
    if r[2] or r[4]:
        all_pairs.append((r[2], r[4], "original"))
    if r[6] or r[7]:
        all_pairs.append((r[6], r[7], "varied"))

dropped = []
kept_count = 0

for en_raw, lun_raw, src in all_pairs:
    en  = _norm(en_raw)  if en_raw  else ""
    lun = _norm(lun_raw) if lun_raw else ""
    lun_lower = lun.lower()

    if not en_raw or not lun_raw:
        reason = "missing value (one side is None/empty)"
    elif not en.strip():
        reason = "english blank after normalisation"
    elif not lun.strip():
        reason = "lunyoro blank after normalisation"
    elif en.lower() == lun_lower:
        reason = "identity pair (EN == LUN)"
    elif len(lun_lower.split()) < 3:
        reason = f"lunyoro too short ({len(lun_lower.split())} token(s), need ≥3)"
    else:
        kept_count += 1
        continue

    dropped.append({"source": src, "english": en, "lunyoro": lun, "reason": reason})

# Check dedup
df_all = pd.DataFrame(
    [{"english": _norm(r[2]) if r[2] else "", "lunyoro": (_norm(r[4]) if r[4] else "").lower()} for r in rows if r[2] and r[4]] +
    [{"english": _norm(r[6]) if r[6] else "", "lunyoro": (_norm(r[7]) if r[7] else "").lower()} for r in rows if r[6] and r[7]]
)
df_all = df_all[df_all["english"].str.strip().astype(bool) & df_all["lunyoro"].str.strip().astype(bool)]
df_all = df_all[df_all["english"].str.lower() != df_all["lunyoro"]]
df_all = df_all[df_all["lunyoro"].str.split().str.len() >= 3]
before_dedup = len(df_all)
df_deduped = df_all.drop_duplicates(subset=["english", "lunyoro"])
dedup_dropped = before_dedup - len(df_deduped)

print(f"Total raw pairs (orig + varied): {len(all_pairs)}")
print(f"Kept (passed all filters):       {len(df_deduped)}")
print(f"Dropped by quality filters:      {len(dropped)}")
print(f"Dropped by deduplication:        {dedup_dropped}")
print(f"  {len(all_pairs)} - {len(dropped)} - {dedup_dropped} = {len(all_pairs) - len(dropped) - dedup_dropped}")
print()

reason_counts = Counter(
    r["reason"] if "too short" not in r["reason"] else "lunyoro too short (< 3 tokens)"
    for r in dropped
)
print("=" * 60)
print("DROP REASONS (summary)")
print("=" * 60)
for reason, count in reason_counts.most_common():
    print(f"  {count:>3}x  {reason}")

# Show short lunyoro examples
short = [r for r in dropped if "too short" in r["reason"]]
if short:
    print(f"\n--- Short Lunyoro examples (all {len(short)}) ---")
    for r in short:
        print(f"  [{r['source']}]")
        print(f"    EN : {r['english'][:80]}")
        print(f"    LUN: \"{r['lunyoro']}\"  ← {r['reason']}")

# Show identity pairs
identity = [r for r in dropped if "identity" in r["reason"]]
if identity:
    print(f"\n--- Identity pairs (EN == LUN) ({len(identity)}) ---")
    for r in identity:
        print(f"  \"{r['english'][:80]}\"")

# Show missing value pairs
missing = [r for r in dropped if "missing" in r["reason"]]
if missing:
    print(f"\n--- Missing value pairs ({len(missing)}) ---")
    for r in missing:
        print(f"  EN: {repr(r['english'][:60])}  LUN: {repr(r['lunyoro'][:60])}")

# Show dedup examples
if dedup_dropped > 0:
    dups = df_all[df_all.duplicated(subset=["english", "lunyoro"], keep="first")]
    print(f"\n--- Duplicate pairs ({dedup_dropped}) ---")
    for _, row in dups.head(10).iterrows():
        print(f"  EN : {row['english'][:70]}")
        print(f"  LUN: {row['lunyoro'][:70]}")
        print()
