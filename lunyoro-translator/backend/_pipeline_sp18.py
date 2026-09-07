"""
_pipeline_sp18.py
=================
Full pipeline for sentence pair 18:
  1. Extract + clean pairs from xlsx
  2. Merge into new_only_train/val (dedup)
  3. Train MarianMT --new-only both directions
  4. Train NLLB     --new-only both directions
  5. Push all 4 models to HF Hub
"""
import re, shutil, subprocess, sys, unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE        = Path(__file__).parent
RAW_DIR     = BASE / "data" / "raw"
CLEANED_DIR = BASE / "data" / "cleaned"
TRAINING_DIR = BASE / "data" / "training"

INPUT_FILE  = next(RAW_DIR.glob("sentence pair 18*"))
CLEAN_FILE  = CLEANED_DIR / "sentence_pair_18_clean.csv"
TRAIN_CSV   = TRAINING_DIR / "new_only_train.csv"
VAL_CSV     = TRAINING_DIR / "new_only_val.csv"

VENV_PY = BASE / "venv312" / "Scripts" / "python.exe"
PY = str(VENV_PY) if VENV_PY.exists() else sys.executable

# ── Step 1: Clean ─────────────────────────────────────────────────────────────
print("=" * 60)
print(f"  SP18 PIPELINE")
print(f"  Input: {INPUT_FILE.name}")
print("=" * 60)

def _norm(text):
    text = str(text).strip()
    text = re.sub(r"\s+", " ", text)
    return (unicodedata.normalize("NFC", text)
            .replace("\u2018","'").replace("\u2019","'")
            .replace("\u201c",'"').replace("\u201d",'"'))

wb = load_workbook(INPUT_FILE)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
print(f"\n[1/5] Cleaning {len(rows)} raw rows...")

orig = pd.DataFrame([(r[2], r[4]) for r in rows if r[2] and r[4]], columns=["english","lunyoro"])
varied = pd.DataFrame([(r[6], r[7]) for r in rows if r[6] and r[7]], columns=["english","lunyoro"])
df = pd.concat([orig, varied], ignore_index=True)

df["english"] = df["english"].apply(_norm)
df["lunyoro"] = df["lunyoro"].apply(_norm).str.lower()
df = df.dropna()
df = df[df["english"].str.strip().astype(bool) & df["lunyoro"].str.strip().astype(bool)]
df = df[df["english"].str.lower() != df["lunyoro"]]
df = df[df["lunyoro"].str.split().str.len() >= 3]
df = df.drop_duplicates(subset=["english","lunyoro"])

CLEANED_DIR.mkdir(exist_ok=True)
df.to_csv(CLEAN_FILE, index=False)
print(f"  Clean pairs: {len(df)} -> {CLEAN_FILE.name}")

# ── Step 2: Merge ─────────────────────────────────────────────────────────────
print(f"\n[2/5] Merging into training set...")
base_train = pd.read_csv(TRAIN_CSV)
base_val   = pd.read_csv(VAL_CSV)
existing   = set(zip(
    pd.concat([base_train, base_val])["english"].str.lower().str.strip(),
    pd.concat([base_train, base_val])["lunyoro"].str.lower().str.strip()
))
new = df[~df.apply(lambda r: (r["english"].lower().strip(), r["lunyoro"].lower().strip()) in existing, axis=1)]
print(f"  Genuinely new: {len(new)}")

if len(new) > 0:
    shutil.copy(TRAIN_CSV, str(TRAIN_CSV)+".bak")
    shutil.copy(VAL_CSV,   str(VAL_CSV)+".bak")
    split = int(len(new) * 0.9)
    new.iloc[:split][["english","lunyoro"]].to_csv(TRAIN_CSV, mode="a", index=False, header=False)
    new.iloc[split:][["english","lunyoro"]].to_csv(VAL_CSV,   mode="a", index=False, header=False)
    t2 = pd.read_csv(TRAIN_CSV); v2 = pd.read_csv(VAL_CSV)
    print(f"  new_only_train: {len(t2)} (+{len(new.iloc[:split])})")
    print(f"  new_only_val  : {len(v2)} (+{len(new.iloc[split:])})")
    # Write timestamped incremental record
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    new.to_csv(CLEANED_DIR / f"incremental_{ts}.csv", index=False)
else:
    print("  Already merged — nothing to do.")

# ── Step 3: MarianMT ──────────────────────────────────────────────────────────
print(f"\n[3/5] Training MarianMT --new-only both directions (5 epochs)...")
subprocess.run([PY, "train_marian.py", "--new-only", "--direction", "both",
                "--epochs", "5", "--batch-size", "32", "--lr", "3e-5"], check=True)

# ── Step 4: NLLB ──────────────────────────────────────────────────────────────
print(f"\n[4/5] Training NLLB --new-only both directions (3 epochs)...")
subprocess.run([PY, "train_nllb.py", "--new-only", "--direction", "both",
                "--epochs", "3", "--batch-size", "16", "--lr", "8e-6", "--fp16"], check=True)

# ── Step 5: Push ──────────────────────────────────────────────────────────────
print(f"\n[5/5] Pushing all 4 models to HF Hub...")
import os
# HF_TOKEN must be set in .env or environment — never hardcoded
for model in ["en2lun", "lun2en", "nllb_en2lun", "nllb_lun2en"]:
    subprocess.run([PY, "push_models.py", "--model", model], check=True)

print("\n" + "=" * 60)
print("  SP18 PIPELINE COMPLETE")
print("=" * 60)
