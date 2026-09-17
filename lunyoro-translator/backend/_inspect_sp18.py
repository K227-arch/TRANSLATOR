import openpyxl, glob
from pathlib import Path

raw = Path("data/raw")
f = list(raw.glob("sentence pair 18*"))[0]
print(f"File: {f.name}")
wb = openpyxl.load_workbook(f)
ws = wb.active
print(f"Sheets: {wb.sheetnames}  |  rows: {ws.max_row}  cols: {ws.max_column}")
print("\nFirst 3 rows:")
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i >= 3: break
    print(f"  {row}")
print("\nLast row:")
rows = list(ws.iter_rows(values_only=True))
print(f"  {rows[-1]}")
